import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from tools import presentation_builder


class PresentationBuilderWithoutDependencyTests(unittest.TestCase):
    def test_invalid_spec_is_rejected_without_powerpoint_dependency(self):
        with self.assertRaises(ValueError):
            presentation_builder.create_deck({"slides": []}, "unused.pptx")

    def test_cli_reports_missing_optional_dependency_without_creating_output(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            source, output = directory / "input.json", directory / "output.pptx"
            source.write_text(json.dumps({"slides": [{"title": "Title", "bullets": []}]}), encoding="utf-8")
            result = subprocess.run([sys.executable, "-S", str(Path(presentation_builder.__file__)), str(source), str(output)],
                                    capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(result.returncode, 2)
            self.assertIn("python-pptx", result.stderr)
            self.assertFalse(output.exists())


@unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx unavailable")
class PresentationBuilderTests(unittest.TestCase):
    def test_shipped_example_produces_a_two_slide_deck(self):
        from pptx import Presentation
        source = Path(presentation_builder.__file__).resolve().parents[1] / 'examples/presentation.json'
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'example.pptx'
            result = subprocess.run([sys.executable, '-B', str(Path(presentation_builder.__file__)),
                                     str(source), str(output)], capture_output=True, text=True, encoding='utf-8')
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(len(Presentation(output).slides), 2)

    def test_figure_slide_embeds_supplied_image_without_distortion(self):
        from PIL import Image
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image_path = root / 'synthetic.png'
            Image.new('RGB', (400, 200), 'blue').save(image_path)
            output = root / 'deck.pptx'
            spec = {'slides': [{'title': 'Example', 'bullets': ['Synthetic figure'],
                                'image': {'path': 'synthetic.png', 'caption': 'Illustrative only'}}]}
            presentation_builder.create_deck(spec, output, base_dir=root)
            slide = Presentation(output).slides[0]
            pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
            self.assertEqual(len(pictures), 1)
            self.assertAlmostEqual(pictures[0].width / pictures[0].height, 2.0, places=4)
            self.assertIn('Illustrative only', [s.text for s in slide.shapes if s.has_text_frame])
            self.assertTrue(image_path.exists())

    def test_cli_resolves_image_relative_to_json_file(self):
        from PIL import Image
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            Image.new('RGB', (60, 120), 'red').save(root / 'figure.png')
            spec = {'slides': [{'title': 'Figure', 'bullets': [], 'image': {'path': 'figure.png'}}]}
            (root / 'spec.json').write_text(json.dumps(spec), encoding='utf-8')
            output = root / 'deck.pptx'
            result = subprocess.run([sys.executable, '-B', str(Path(presentation_builder.__file__)),
                                     str(root / 'spec.json'), str(output)], cwd=root.parent,
                                    capture_output=True, text=True, encoding='utf-8')
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(any(s.shape_type == 13 for s in Presentation(output).slides[0].shapes))

    def test_invalid_or_missing_figure_never_creates_output(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            invalid = root / 'bad.png'
            invalid.write_bytes(b'not an image')
            for filename in ('bad.png', 'missing.png'):
                output = root / (filename + '.pptx')
                spec = {'slides': [{'title': 'Title', 'bullets': [], 'image': {'path': filename}}]}
                with self.subTest(filename=filename), self.assertRaises((OSError, ValueError)):
                    presentation_builder.create_deck(spec, output, base_dir=root)
                self.assertFalse(output.exists())

    def test_oversized_image_header_is_reported_without_creating_output(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'large.png').write_bytes(b'header')
            output = root / 'deck.pptx'
            spec = {'slides': [{'title': 'Title', 'bullets': [], 'image': {'path': 'large.png'}}]}
            with mock.patch.object(Image, 'open', side_effect=Image.DecompressionBombError('too large')):
                with self.assertRaisesRegex(ValueError, 'image'):
                    presentation_builder.create_deck(spec, output, base_dir=root)
            self.assertFalse(output.exists())

    def test_failed_serialization_leaves_no_output_or_temporary_files(self):
        from pptx.presentation import Presentation as PresentationClass

        def fail_save(self, destination):
            destination.write(b"partial")
            raise OSError("simulated save failure")

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "deck.pptx"
            with mock.patch.object(PresentationClass, "save", fail_save):
                with self.assertRaisesRegex(OSError, "simulated save failure"):
                    presentation_builder.create_deck({"slides": [{"title": "Title", "bullets": []}]}, path)
            self.assertEqual(list(Path(directory).iterdir()), [])

    def test_saved_deck_reopens_with_editable_text_and_shapes(self):
        from pptx import Presentation
        spec = {"slides": [{"title": "Research question", "bullets": ["First point", "Second point"]}]}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "deck.pptx"
            presentation_builder.create_deck(spec, path)
            deck = Presentation(path)
            self.assertEqual(len(deck.slides), 1)
            texts = [shape.text for shape in deck.slides[0].shapes if shape.has_text_frame]
            self.assertIn("Research question", texts)
            self.assertIn("• First point\n• Second point", texts)
            self.assertGreaterEqual(len(deck.slides[0].shapes), 3)

    def test_rejects_invalid_inputs_before_writing_and_does_not_overwrite(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "deck.pptx"
            for spec in ({"slides": []}, {"slides": [{"title": "x", "bullets": [1]}]},
                         {"slides": [{"title": "x", "bullets": ["x"] * 7}]}):
                with self.subTest(spec=spec), self.assertRaises(ValueError):
                    presentation_builder.create_deck(spec, path)
                self.assertFalse(path.exists())
            path.write_bytes(b"existing")
            with self.assertRaises(FileExistsError):
                presentation_builder.create_deck({"slides": [{"title": "Title", "bullets": []}]}, path)
            self.assertEqual(path.read_bytes(), b"existing")


if __name__ == "__main__":
    unittest.main()
