"""Create a simple editable research deck from caller-supplied JSON content.

Requires the optional python-pptx package. Visual inspection is still needed.
"""
import argparse
import io
import json
import os
import sys
import tempfile
from pathlib import Path


def _validate(spec):
    if not isinstance(spec, dict) or set(spec) != {"slides"}:
        raise ValueError("input must be an object containing only slides")
    slides = spec["slides"]
    if not isinstance(slides, list) or not 1 <= len(slides) <= 100:
        raise ValueError("slides must be a list of 1 to 100 slides")
    for i, slide in enumerate(slides, 1):
        if not isinstance(slide, dict) or not {"title", "bullets"} <= set(slide) or set(slide) - {"title", "bullets", "image"}:
            raise ValueError(f"slide {i} must contain title and bullets, with optional image")
        title, bullets = slide["title"], slide["bullets"]
        if not isinstance(title, str) or not title.strip() or len(title) > 90 or any(c in title for c in "\r\n"):
            raise ValueError(f"slide {i} title must be a single nonempty line of at most 90 characters")
        if not isinstance(bullets, list) or len(bullets) > 6:
            raise ValueError(f"slide {i} bullets must be a list of at most 6 lines")
        for line in bullets:
            if not isinstance(line, str) or not line.strip() or len(line) > 115 or any(c in line for c in "\r\n"):
                raise ValueError(f"slide {i} bullet must be a single nonempty line of at most 115 characters")
        if 'image' in slide:
            image = slide['image']
            if (not isinstance(image, dict) or not set(image) <= {'path', 'caption'} or
                    not isinstance(image.get('path'), str) or not image['path'].strip() or
                    len(image['path']) > 1000 or '\x00' in image['path'] or
                    ('caption' in image and (not isinstance(image['caption'], str) or len(image['caption']) > 150 or
                                             any(c in image['caption'] for c in '\r\n')))):
                raise ValueError(f"slide {i} image needs a local path and optional single-line caption")


def create_deck(spec, output, *, base_dir=None):
    """Create a new PPTX; refuse overwrite and return its path."""
    _validate(spec)
    output = Path(output)
    if output.suffix.lower() != ".pptx":
        raise ValueError("output must have a .pptx extension")
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError("Install the optional python-pptx package to create a deck") from exc

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    base_dir = Path(base_dir) if base_dir is not None else Path.cwd()
    for item in spec["slides"]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(.7), Inches(.5), Inches(11.9), Inches(1.1))
        title.text_frame.word_wrap = True
        title.text_frame.text = item["title"]
        for run in title.text_frame.paragraphs[0].runs:
            run.font.size, run.font.bold = Pt(30), True
            run.font.color.rgb = RGBColor(25, 50, 75)
        line = slide.shapes.add_shape(1, Inches(.7), Inches(1.75), Inches(11.9), Inches(.04))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(40, 100, 140)
        line.line.fill.background()
        body = slide.shapes.add_textbox(Inches(.8), Inches(2), Inches(5.5 if 'image' in item else 11.7), Inches(4.8))
        frame = body.text_frame
        frame.word_wrap = True
        for i, value in enumerate(item["bullets"]):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = "• " + value
            paragraph.space_after = Pt(18)
            for run in paragraph.runs:
                run.font.size = Pt(22)
                run.font.color.rgb = RGBColor(35, 35, 35)
        if 'image' in item:
            image = item['image']
            image_path = base_dir / image['path']
            with image_path.open('rb') as stream:
                data = stream.read(20_000_001)
            if len(data) > 20_000_000:
                raise ValueError('image exceeds 20 MB')
            try:
                with Image.open(io.BytesIO(data)) as picture:
                    width_px, height_px = picture.size
                    if picture.format not in ('PNG', 'JPEG') or width_px * height_px > 50_000_000 or not width_px or not height_px:
                        raise ValueError('image must be PNG or JPEG and at most 50 megapixels')
            except Image.DecompressionBombError as exc:
                raise ValueError('image exceeds the supported dimensions') from exc
            max_w, max_h = Inches(5.2), Inches(3.7)
            scale = min(max_w / width_px, max_h / height_px)
            width, height = int(width_px * scale), int(height_px * scale)
            left, top = Inches(7.1) + (max_w - width) // 2, Inches(2.05) + (max_h - height) // 2
            slide.shapes.add_picture(io.BytesIO(data), left, top, width=width, height=height)
            if image.get('caption'):
                caption = slide.shapes.add_textbox(Inches(7.1), Inches(5.85), max_w, Inches(.65))
                caption.text_frame.word_wrap = True
                caption.text_frame.text = image['caption']
                for run in caption.text_frame.paragraphs[0].runs:
                    run.font.size = Pt(14)
    # Write a complete file first, then publish it without replacing any existing
    # path, including one created concurrently. Both paths share a filesystem.
    fd, temporary = tempfile.mkstemp(prefix=".presentation-", suffix=".tmp", dir=output.parent)
    try:
        with os.fdopen(fd, "wb") as destination:
            deck.save(destination)
        os.link(temporary, output)
    finally:
        os.unlink(temporary)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", help="UTF-8 JSON specification")
    parser.add_argument("output", help="new .pptx path; existing files are never overwritten")
    args = parser.parse_args()
    try:
        spec = json.loads(Path(args.input).read_text(encoding="utf-8"))
        create_deck(spec, args.output, base_dir=Path(args.input).resolve().parent)
    except (OSError, ValueError, RuntimeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
